from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = r"C:\Users\HP\projects\permitos\docs\EstatePermit_KCMO_Five_Step_Implementation_Plan.docx"
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 98, 108)


def font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    font(p.add_run(text))


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 14, 7), ("Heading 3", 12, DARK, 10, 5)
]:
    s = styles[name]; s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = color
    s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after); s.paragraph_format.keep_with_next = True

# Running furniture
h = sec.header.paragraphs[0]
h.text = "ESTATEPERMIT  /  KANSAS CITY MVP"
h.alignment = WD_ALIGN_PARAGRAPH.LEFT
font(h.runs[0], 9, True, MUTED)
f = sec.footer.paragraphs[0]
f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(f.add_run("Implementation plan  |  July 2026"), 9, False, MUTED)

# Title block: memo masthead without a border.
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
font(p.add_run("IMPLEMENTATION PLAN"), 10, True, BLUE)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
font(p.add_run("EstatePermit: Kansas City Permit Management MVP"), 24, True, DARK)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(16)
font(p.add_run("A five-step roadmap from permit recommendation to complete operational tracking"), 13, False, MUTED)

for label, value in [("Objective", "Support one construction project with a coordinated bundle of independently managed permits."),
                     ("Initial jurisdiction", "Kansas City, Missouri (KCMO)"),
                     ("MVP boundary", "Steps 1-3 create the usable MVP; Steps 4-5 complete the operating lifecycle.")]:
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    font(p.add_run(label + ": "), 11, True, DARK); font(p.add_run(value), 11)

doc.add_heading("Product model", level=1)
p = doc.add_paragraph()
font(p.add_run("Core principle: "), 11, True, DARK)
font(p.add_run("one project contains one coordinated permit bundle, and every permit remains independently manageable."))

table = doc.add_table(rows=2, cols=3)
table.autofit = False
widths = [Inches(1.65), Inches(2.45), Inches(2.40)]
for row in table.rows:
    for i, cell in enumerate(row.cells): cell.width = widths[i]
headers = ["Project", "Permit bundle", "Permit records"]
values = ["Address, development type and construction scope", "All required, likely and optional approvals", "Status, authority, documents, dependencies and inspections"]
for i, text in enumerate(headers):
    shade(table.cell(0, i), "E8EEF5"); font(table.cell(0, i).paragraphs[0].add_run(text), 10, True, DARK)
for i, text in enumerate(values): font(table.cell(1, i).paragraphs[0].add_run(text), 10)

steps = [
    ("1. Upgrade intake and add KCMO permit knowledge",
     "Give the recommendation engine enough structured information to determine the complete permit bundle.",
     ["Rename Project type to Development type; retain options such as single-family, multifamily, commercial tenant improvement, mixed-use and industrial.",
      "Add structured scope questions for construction activity, structural work, electrical, plumbing, HVAC, fire alarm, sprinklers, signs, demolition, change of use, grading, right-of-way, solar and water/sewer work.",
      "Create the kansas_city_mo knowledge pack with permit types, triggers, authorities, dependencies, required documents and citations.",
      "Confirm municipal jurisdiction from the address and distinguish KCMO from Kansas City, Kansas and surrounding municipalities."],
     "EstatePermit can generate an explainable KCMO permit bundle from normal project intake."),
    ("2. Create persistent permit records",
     "Convert temporary analysis recommendations into operational records that survive reruns and user edits.",
     ["Add a ProjectPermit database entity related to Project.",
      "Store permit type, agency, requirement classification, lifecycle status, parent/dependency relationships, CompassKC number, owner, contractor, fees, dates, blockers and next action.",
      "Keep requirement status separate from lifecycle status—for example Required versus Submitted.",
      "Allow users to confirm, add or dismiss recommendations without a later analysis overwriting managed records."],
     "Every recommended permit can become a durable, independently managed record."),
    ("3. Build the Permit Bundle workspace",
     "Make multi-permit coordination the primary day-to-day experience.",
     ["Add a permanent Permits tab to each project.",
      "Show permit name, authority, requirement status, lifecycle status, responsible party, dependencies, missing documents and next action.",
      "Provide a detail view with Overview, Documents, Corrections, Inspections and Activity sections.",
      "Visualize blockers such as an electrical permit waiting on the parent commercial building permit."],
     "Users can manage the entire bundle while updating each permit independently."),
    ("4. Add document readiness and correction management",
     "Prevent incomplete submissions and structure the review/resubmission loop.",
     ["Keep each uploaded file once at project level and link it to one or more permits through PermitDocumentLink.",
      "Generate KCMO-specific document checklists and track required, received, reviewed and accepted states.",
      "Add document versions, reviewer comments, responsible owners, due dates, resolutions and resubmission rounds.",
      "Generate a clear readiness result for every permit and for the overall project."],
     "Teams know exactly what is missing and who must resolve each correction."),
    ("5. Add inspections, completion and portal operations",
     "Cover the lifecycle from issued permit through final approval and occupancy.",
     ["Add permit-level inspections with type, requested/scheduled date, result, notes, failed items and reinspection status.",
      "Track fees, payment state, issuance, expiration, finals and Certificate of Occupancy.",
      "Add CompassKC deep links, manual status-check reminders and a complete activity history.",
      "Treat automated CompassKC synchronization as a later enhancement unless a stable supported integration is available."],
     "EstatePermit supports the complete operational lifecycle through final inspection."),
]

for title, goal, actions, outcome in steps:
    doc.add_heading(title, level=1)
    p = doc.add_paragraph(); font(p.add_run("Goal: "), 11, True, DARK); font(p.add_run(goal))
    doc.add_heading("Implementation work", level=2)
    for a in actions: add_bullet(doc, a)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
    font(p.add_run("Completion outcome: "), 11, True, DARK); font(p.add_run(outcome))

doc.add_heading("Recommended delivery gates", level=1)
for text in [
    "Gate 1 — Detection works: representative KCMO projects produce correct, explainable permit bundles.",
    "Gate 2 — Records persist: confirmed permits survive analysis reruns and retain user-entered status.",
    "Gate 3 — MVP usable: a contractor can create a project, confirm permits, assign work and see blockers.",
    "Gate 4 — Submission ready: each permit has an auditable document checklist and correction workflow.",
    "Gate 5 — Lifecycle complete: issued permits, inspections, finals and occupancy can all be tracked."
]: add_bullet(doc, text)

p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
font(p.add_run("Recommended release boundary: "), 11, True, DARK)
font(p.add_run("Release the first customer-facing MVP after Step 3, then use live customer projects to prioritize the document and inspection features in Steps 4 and 5."))

doc.save(OUT)
print(OUT)
