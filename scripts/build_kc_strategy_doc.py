from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r"C:\Users\HP\projects\EstatePermit\docs\EstatePermit_Kansas_City_MVP_and_Go_to_Market.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "17324D"; BLUE = "246B9E"; PALE = "E8F1F7"; LIGHT = "F4F6F8"
GREEN = "176B4D"; GOLD = "8A6300"; RED = "9A2C2C"; GRAY = "5D6770"; WHITE = "FFFFFF"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.78); sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(0.82); sec.right_margin = Inches(0.82)
sec.header_distance = Inches(0.35); sec.footer_distance = Inches(0.35)

def font(run, size=10.5, bold=False, color="222222", italic=False):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"; normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.12
for name, size, color, before, after in [
    ("Heading 1", 17, NAVY, 15, 7), ("Heading 2", 13.5, BLUE, 11, 5),
    ("Heading 3", 11.5, NAVY, 8, 3)]:
    s=styles[name]; s.font.name="Aptos Display"; s.font.size=Pt(size); s.font.bold=True
    s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before)
    s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True

for st in ["List Bullet", "List Number"]:
    styles[st].font.name="Aptos"; styles[st].font.size=Pt(10.5)
    styles[st].paragraph_format.space_after=Pt(3); styles[st].paragraph_format.line_spacing=1.1

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=90, start=110, bottom=90, end=110):
    tc=cell._tc.get_or_add_tcPr(); mar=tc.first_child_found_in("w:tcMar")
    if mar is None: mar=OxmlElement("w:tcMar"); tc.append(mar)
    for edge,val in [("top",top),("start",start),("bottom",bottom),("end",end)]:
        el=mar.find(qn(f"w:{edge}"))
        if el is None: el=OxmlElement(f"w:{edge}"); mar.append(el)
        el.set(qn("w:w"),str(val)); el.set(qn("w:type"),"dxa")

def set_cell_text(cell, text, bold=False, color="222222", size=9.4):
    cell.text=""; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
    r=p.add_run(str(text)); font(r,size,bold,color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for i,h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, WHITE, 9.2); shade(t.rows[0].cells[i], NAVY)
        if widths: t.rows[0].cells[i].width=Inches(widths[i])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            set_cell_text(cells[i], val, False, "222222", 9.1)
            if widths: cells[i].width=Inches(widths[i])
            if ri%2: shade(cells[i], "F7F9FA")
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def p(text="", bold_lead=None, color="222222", size=10.5, align=None, after=5):
    para=doc.add_paragraph(); para.paragraph_format.space_after=Pt(after)
    if align is not None: para.alignment=align
    if bold_lead and text.startswith(bold_lead):
        r=para.add_run(bold_lead); font(r,size,True,color)
        r=para.add_run(text[len(bold_lead):]); font(r,size,False,color)
    else:
        r=para.add_run(text); font(r,size,False,color)
    return para

def bullet(text, level=0):
    para=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2")
    para.paragraph_format.left_indent=Inches(0.25+0.2*level); para.paragraph_format.first_line_indent=Inches(-0.14)
    r=para.add_run(text); font(r,10.3)
    return para

def numbered(text):
    para=doc.add_paragraph(style="List Number"); r=para.add_run(text); font(r,10.3); return para

def callout(label, text, fill=PALE, accent=BLUE):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    c=t.cell(0,0); c.width=Inches(6.75); shade(c,fill); margins(c,150,180,150,180)
    c.text=""; para=c.paragraphs[0]; para.paragraph_format.space_after=Pt(0)
    r=para.add_run(label.upper()+"  "); font(r,9.3,True,accent)
    r=para.add_run(text); font(r,10.4,False,NAVY)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def page_break(): doc.add_page_break()

# Header/footer
hdr=sec.header.paragraphs[0]; hdr.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=hdr.add_run("EstatePermit  |  KANSAS CITY MVP STRATEGY"); font(r,8.5,True,GRAY)
ftr=sec.footer.paragraphs[0]; ftr.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=ftr.add_run("Working strategy - July 2026"); font(r,8.2,False,GRAY)

# Cover
doc.add_paragraph().paragraph_format.space_after=Pt(55)
p("EstatePermit", color=BLUE, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
title=p("Kansas City MVP and Go-to-Market Playbook", color=NAVY, size=28, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
title.runs[0].bold=True
p("What to build, how to validate it, and who to approach first", color=GRAY, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=28)
callout("Core decision", "Build a credible permit-readiness MVP for Kansas City, Missouri; validate it with Kansas State University reviewers and friendly practitioners; then run free concierge pilots with small and mid-size contractors before expanding across the metro or to St. Louis.")
p("Prepared for the EstatePermit founding team", color=GRAY, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
p("Research date: July 11, 2026", color=GRAY, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
page_break()

doc.add_heading("Executive direction", level=1)
p("The proposed sequence is sound: first produce a stable MVP, test it with trusted reviewers who understand design and construction, then begin structured discovery and free pilots with Kansas City companies. The product should not initially claim to replace architects, engineers, code officials, or experienced permit expediters. It should make the administrative and preflight work faster, clearer, and more traceable.")
callout("The first commercial promise", "Give us the project facts and documents. EstatePermit identifies the likely permit path, organizes required materials, flags missing or inconsistent information, produces a reviewable readiness report, and keeps every action visible. A qualified human remains responsible for technical judgments and final filing.", "EAF5EF", GREEN)
doc.add_heading("Scope decision", level=2)
table(["Now", "Next", "Later"], [[
    "Kansas City, Missouri; one or two repeatable project types; readiness and packet preparation",
    "Kansas City metro, including KCK/Wyandotte and Johnson County jurisdictions",
    "Greater Kansas and St. Louis after repeatable paid demand"
]], [2.25,2.25,2.25])
doc.add_heading("What success looks like before broad outreach", level=2)
for x in [
    "A complete demo project can be created, analyzed, reviewed, exported, and explained without developer intervention.",
    "Three to five trusted reviewers can comment on the output and identify missing requirements or unsafe claims.",
    "At least ten historical KCMO permits can be replayed through the workflow and compared with actual outcomes.",
    "Every output distinguishes sourced fact, system inference, user-provided information, and professional judgment.",
    "The team can offer a free pilot with clear scope, privacy terms, and a human approval gate."
]: bullet(x)

doc.add_heading("Part I - Product and MVP specification", level=1)
doc.add_heading("1. What EstatePermit already has", level=2)
p("The repository is more than a concept. It contains a working project-first React interface, FastAPI persistence, file handling, jurisdiction packs, multi-module analysis, report packaging, and an audit-oriented approval workflow.")
table(["Current capability", "Observed implementation", "MVP assessment"], [
    ["Project workspace", "Create/edit/delete projects; address, type, jurisdiction, area", "Keep"],
    ["Files", "Upload, classify, label, assign to analysis modules, remove", "Keep and harden"],
    ["Rules", "Built-in and custom rules; suggested rules; per-module runs", "Keep, simplify for users"],
    ["Analysis", "Zoning, building, fire/site modules; pass/fail/warn; citations", "Core differentiator"],
    ["Permit package", "Required permits, documents, fees, timeline, filing sequence", "Convert into permit cards"],
    ["Human gate", "Approve for filing; audit hash; approval metadata", "Keep but rename carefully"],
    ["RFI demonstration", "Simulated city RFI and generated response draft", "Useful demo; not yet operational"],
    ["Reporting", "Downloadable PDF analysis report", "Keep and improve"],
    ["Portfolio dashboard", "Search and filter projects; readiness summaries", "Keep; add next action and aging"]
], [1.45,3.45,1.85])
callout("Current technical caveat", "The frontend build could not be independently completed in this review because the bundled package manager blocked ignored dependency build scripts, and the bundled Python runtime did not include pytest. Treat clean build and automated test execution as an MVP release gate, not as evidence that the product is presently broken.", "FFF6E0", GOLD)

doc.add_heading("2. MVP feature set", level=2)
p("The following is the minimum credible set for a company demonstration. Items marked Essential should be complete before cold outreach. Items marked Pilot can be operated manually behind the interface during early free engagements.")
table(["Feature", "What the user gets", "Level"], [
    ["Guided project intake", "Address, parcel, scope, valuation, occupancy/use, trades, contacts, target start date", "Essential"],
    ["Document intake", "Drag-and-drop files, type labels, versions, preview, missing-file warnings", "Essential"],
    ["Project fact sheet", "Normalized facts extracted from forms/files, each with source and confidence", "Essential"],
    ["Permit detection", "Likely permit cards with a plain-language reason each applies", "Essential"],
    ["Requirement checklist", "Required, received, missing, uncertain, and not-applicable items", "Essential"],
    ["Readiness review", "Administrative completeness plus rule-based risk findings; no claim of professional approval", "Essential"],
    ["Evidence drawer", "Official source, link, effective/update date, excerpt, and verification date", "Essential"],
    ["Issue workflow", "Severity, owner, due date, recommended action, resolved state", "Essential"],
    ["Human review gate", "Reviewer can approve, request changes, or mark expert review required", "Essential"],
    ["Exportable packet", "Branded readiness report, checklist, source appendix, and document index", "Essential"],
    ["Comments and feedback", "Reviewers comment on project, permit, requirement, finding, or document", "Essential"],
    ["Activity history", "Who changed, reviewed, approved, or exported what and when", "Essential"],
    ["KCMO status tracker", "Manual permit number, CompassKC link, current status, next action, important dates", "Pilot"],
    ["Correction log", "Reviewer comment, assigned responder, revised document, resubmission status", "Pilot"],
    ["Email/reminder assistance", "Templates and reminders; no autonomous external sending", "Pilot"],
    ["Inspection checklist", "Required inspections, readiness, request date, result, reinspection", "Pilot after issuance"]
], [1.55,4.15,1.05])

doc.add_heading("3. What the dashboard should show", level=2)
doc.add_heading("Portfolio dashboard", level=3)
for x in [
    "Projects by lifecycle state and readiness, not only analysis count.",
    "Next action, assigned owner, due date, days waiting, and last city/customer update.",
    "Filters for project type, permit type, contractor/customer, jurisdiction, risk, and aging.",
    "Attention queue: missing documents, unresolved high-risk findings, expiring permits, and overdue corrections.",
    "Simple outcome metrics: time to ready, time to submit, corrections per permit, and time to issuance."
]: bullet(x)
doc.add_heading("Project dashboard", level=3)
for x in [
    "Header: project, address, jurisdiction, customer, responsible reviewer, lifecycle state, readiness, and next action.",
    "Project facts: a source-linked, editable summary of scope, use, valuation, area, trades, and key dates.",
    "Permit cards: applicability reason, agency, status, checklist completion, risk, fees, timeline, and owner.",
    "Documents: latest version, document type, matched requirements, review state, and missing-document map.",
    "Issues: blockers, contradictions, assignments, deadlines, comments, and evidence.",
    "Timeline: intake, analysis, review decisions, submissions, corrections, inspections, and closeout events.",
    "Actions: request information, rerun review, export packet, copy email draft, record submission, request approval."
]: bullet(x)

doc.add_heading("4. Features to defer", level=2)
table(["Do not build for the first MVP", "Reason"], [
    ["Nationwide jurisdiction coverage", "Creates shallow, untestable knowledge and destroys focus"],
    ["Autonomous portal filing", "High reliability, security, authorization, and liability burden"],
    ["Automatic fee payment", "Financial and authorization risk"],
    ["Claims of code compliance or guaranteed approval", "Requires professional judgment and creates liability"],
    ["Full CRM/accounting/project-management suite", "Customers already have systems; integration comes later"],
    ["Complex generative plan review", "Hard to validate without disciplined expert and dataset support"],
    ["St. Louis or multiple Kansas jurisdictions", "Each jurisdiction is a separate rules and operations product"]
], [3.0,3.75])

doc.add_heading("5. MVP acceptance gates", level=2)
for x in [
    "Reliability: clean install/build, automated tests runnable, no silent analysis failure, recoverable job status.",
    "Accuracy: every KCMO requirement has an official source and date; uncertain cases are labeled uncertain.",
    "Safety: no output represents itself as architecture, engineering, legal advice, or municipal approval.",
    "Usability: a first-time reviewer can complete the demo without verbal coaching from a developer.",
    "Evidence: every finding and permit recommendation can be traced to project data and a source.",
    "Feedback: comments can be captured at the exact object that caused confusion.",
    "Privacy: customer files are access-controlled, removable, and excluded from model training by default.",
    "Demo: one complete and one incomplete KCMO example show obvious before/after value."
]: bullet(x)

page_break()
doc.add_heading("Part II - Kansas City market and permitting playbook", level=1)
doc.add_heading("6. KCMO permitting process", level=2)
p("Kansas City, Missouri administers permits, contractor licensing, plan review, inspections, and certificates through City Planning and Development. CompassKC is the principal online system for submitting plans and permits, checking records, and requesting inspections.")
table(["Stage", "Operational reality", "EstatePermit opportunity"], [
    ["1. Scope and zoning", "Confirm work, property, use, zoning, and whether separate approvals apply", "Guided intake; parcel/source links; uncertainty flags"],
    ["2. Licensed party", "Many structural and trade permits require appropriately licensed KCMO contractors", "License field and validation reminder; never impersonate applicant"],
    ["3. Application", "Applicant creates the relevant CompassKC record and enters project information", "Field checklist; reusable fact sheet; reviewed copy/paste packet"],
    ["4. Plans/documents", "Upload scope, drawings, applications, and permit-specific supporting materials", "Document index, naming/version checks, missing-item map"],
    ["5. Plan review", "Departments review; comments and revisions may follow", "Comment intake, assignments, due dates, response matrix"],
    ["6. Fees and issuance", "Fees must be paid before a permit is valid; trade sequencing may depend on building approval", "Fee status, dependency map, issuance checklist"],
    ["7. Construction/inspections", "Inspections are requested through CompassKC, email, or phone and must precede later stages", "Inspection plan, reminders, results, reinspection tasks"],
    ["8. Completion", "Final approvals and, where applicable, Certificate of Occupancy/TCO", "Closeout checklist, document archive, completion evidence"]
], [1.25,3.15,2.35])
doc.add_heading("KCMO rules that should become product logic", level=3)
for x in [
    "A permit generally precedes construction, alteration, repair, demolition, occupancy change, or regulated electrical, gas, mechanical, and plumbing work.",
    "Construction begins only after fees are paid and the permit is issued.",
    "A permit can expire when work does not begin or an inspection is not obtained within the relevant 180-day period.",
    "Commercial/multifamily trade permits commonly depend on approved building plans and building permit issuance.",
    "Inspection requests have operational cutoffs; KCMO currently states 4:00 p.m. for most next-day requests, subject to daily capacity.",
    "Inspection activity and reports can be tracked through CompassKC.",
    "New or remodeled structures cannot be occupied until the required certificate or written approval is issued.",
    "KCMO information bulletins are critical sources and carry different update dates; source freshness must be recorded."
]: bullet(x)
callout("Important boundary", "EstatePermit should help the licensed applicant prepare, organize, and track work. It should not submit under another party's license, certify plans, or make technical representations that require a licensed professional.", "FBEAEA", RED)

doc.add_heading("7. Recommended first KCMO wedge", level=2)
p("The best first wedge is not a complex commercial tenant improvement. It is a repeated administrative workflow performed by licensed contractors with enough permit volume to feel the pain and enough internal simplicity to try a free pilot.")
table(["Candidate", "Attractiveness", "Risk", "Decision"], [
    ["Residential mechanical replacement", "High frequency; structured data; direct contractor buyer", "Exceptions and equipment/scope differences", "Best discovery candidate"],
    ["Residential electrical service/panel", "Repeated workflow and clear inspection path", "Safety-sensitive; licensing and utility coordination", "Second candidate"],
    ["Residential plumbing replacement", "Frequent and document-light", "Water service vs plumbing distinctions", "Second candidate"],
    ["Residential remodel/basement", "Visible document-readiness value", "Plans, multiple trades, more review variation", "Good professor/friend demo"],
    ["Commercial tenant improvement", "Higher value and richer workflow", "More departments, professional plans, longer cycles", "Phase after first pilots"],
    ["Zoning/development cases", "High pain and high willingness to pay", "Expert-heavy, discretionary, slow", "Defer"]
], [1.7,2.3,1.8,1.0])

doc.add_heading("8. Competitive and customer landscape", level=2)
p("KCMO's official open permit dataset includes contractor company name, trade, permit type, status, dates, address, project description, estimated cost, and a CompassKC record link. The latest published dataset identified in this review was updated through May 9, 2025. It is valuable for market sizing and outreach, but names require normalization because the same company may appear under multiple spellings.")
doc.add_heading("Direct and adjacent competition", level=3)
table(["Type", "Example or signal", "How to respond"], [
    ["Local permit service", "Permit Service Inc. appears under duplicate spellings with more than 3,300 permits in the 2024-May 2025 query", "Treat as an entrenched operator and learning benchmark, not an early sales target"],
    ["Internal contractor staff", "Large HVAC/plumbing/electrical firms show hundreds of permits", "Sell workflow visibility or overflow later; interview operations staff now"],
    ["Architects/engineers", "AIA KC lists more than 80 member firms", "Partners and expert channels, not enemies"],
    ["General software", "CompassKC is the official transaction portal", "Complement it; do not try to replace the municipal system"],
    ["National platform", "PermitFlow offers broad managed lifecycle operations", "Differentiate through KCMO depth, evidence, and small-firm service"]
], [1.45,3.35,1.95])

doc.add_heading("Initial activity-ranked prospect universe", level=3)
p("The following names are not endorsements. Counts are raw permit rows from an official KCMO open-data query for issued dates from January 1, 2024 through May 9, 2025 and may include duplicates, related entities, or data-entry variation.", color=GRAY, size=9.3)
table(["Company name in dataset", "Raw permits", "Outreach posture"], [
    ["Anthony Plumbing, Heating, Cooling & Electric", "789 (+ related name entries)", "Large benchmark/interview; later enterprise"],
    ["A.B. May Company, Inc.", "539", "Large benchmark/interview; later enterprise"],
    ["Total Home Service", "332", "Mid/large discovery target"],
    ["Hamilton Plumbing Heating AC Rooter", "310 (+ variants)", "Large benchmark/interview"],
    ["Mike Bryant Heating & Cooling, LLC", "198", "Strong discovery target"],
    ["C.M. Mose & Son", "159", "Strong discovery target"],
    ["1KCConstruction", "142", "Strong pilot/discovery target"],
    ["KASA Electric LLC", "137", "Strong pilot/discovery target"],
    ["Arrow Circle Electric, Inc.", "122", "Strong pilot/discovery target"],
    ["Allen's Electric, Incorporated", "120", "Strong pilot/discovery target"],
    ["Romans Remodeling / Romans Remodeling LLC", "200 combined raw", "Strong pilot after deduplication"],
    ["McDaniel Furnace & Sheet Metal, Inc.", "105", "Strong pilot/discovery target"],
    ["Buckner's Heating & Cooling", "94", "Strong pilot/discovery target"],
    ["Mister Sparky Kansas City", "89", "Discovery target; likely systemized buyer"],
    ["Teague Electric Construction, Inc.", "87", "Strong pilot/discovery target"],
    ["Hometown Plumbing KC LLC", "85", "Strong pilot/discovery target"],
    ["Scott Hagen Electric", "83", "Strong pilot/discovery target"],
    ["B & L Plumbing Service, Inc.", "79", "Strong pilot/discovery target"],
    ["Lancaster Brothers Heating & Cooling, LLC", "76", "Strong pilot/discovery target"],
    ["A & S Mechanical", "74", "Strong pilot/discovery target"]
], [3.8,0.8,2.15])

doc.add_heading("How to create the complete target database", level=3)
for x in [
    "Export the KCMO Permits CPD dataset monthly through the Socrata API.",
    "Normalize company names, punctuation, DBAs, license numbers, addresses, and related entities.",
    "Calculate permit volume, permit mix, average project value, correction/aging signals where available, and recency.",
    "Join against KCMO's active licensed-contractor lookup and remove inactive or unverified entities.",
    "Enrich only from public business sites and association directories; record source and verification date.",
    "Score prospects: repeated permits + small/mid-size operation + reachable owner/operations lead + administrative pain.",
    "Do not mass-email the entire dataset. Start with 20 highly relevant companies and personalized evidence-based outreach."
]: numbered(x)

doc.add_heading("9. Validation and outreach", level=2)
doc.add_heading("Can outreach start with an MVP?", level=3)
p("Yes. You do not need a finished commercial platform to request a short interview. You do need a credible demo, a precise research purpose, and honesty that the product is in pilot development. Discovery outreach should begin after the core demo is stable; warm interviews with professors, friends, and practitioners can begin immediately.")
callout("Do not lead with a sale", "Lead with the workflow you are studying and the specific evidence that makes the company relevant: for example, repeated KCMO mechanical or electrical permit activity. Ask to understand their current process and offer a free pilot only after confirming a real pain.")
doc.add_heading("Minimum material to have in hand", level=3)
for x in [
    "A two-minute live demo using a realistic KCMO project.",
    "One-page explanation of the problem, product boundary, and free pilot.",
    "Sample readiness report and permit checklist.",
    "Short data/privacy statement and deletion policy.",
    "Pilot scope: one project, no filing representation, customer approval required, no guarantee of approval.",
    "A structured 10-minute interview guide and permission to take notes.",
    "A visible feedback mechanism inside the product or a linked review form."
]: bullet(x)

doc.add_heading("Suggested outreach message", level=3)
callout("Email or LinkedIn", "Hi [Name] - we are a Kansas-based computer science team building EstatePermit, a tool that organizes KCMO permit requirements, project documents, and correction follow-up for small contractors. We are not selling a finished platform or claiming to replace licensed professionals. We are interviewing a small number of companies that regularly work through CompassKC. Would you be open to a 10-minute conversation about how your team currently prepares and tracks permits? We can also run one project through our pilot at no cost in exchange for candid feedback.", "F4F6F8", NAVY)
doc.add_heading("Interview questions", level=3)
for x in [
    "Who prepares and submits permits today, and how much time does that person spend per permit?",
    "Show us the last permit that took longer than expected. Where did the delay begin?",
    "Which information or document is most often missing at intake?",
    "How do you track CompassKC status, reviewer comments, inspections, and expiration risk?",
    "What gets copied between your CRM, estimates, contracts, forms, and the city portal?",
    "Which decision requires an owner, supervisor, architect, or engineer?",
    "What would make you refuse to trust a permitting product?",
    "If we handled the administrative preparation and tracking, what outcome would justify paying?"
]: numbered(x)

doc.add_heading("10. Overall execution sequence", level=2)
table(["Stage", "Primary work", "Exit condition"], [
    ["A. Credible MVP", "Finish essential features, clean build/tests, KCMO sources, two demo projects", "Demo works without developer rescue"],
    ["B. Trusted review", "K-State professors, landscape/design contacts, contractors, permit professional", "Major gaps logged; unsafe claims removed"],
    ["C. Historical validation", "Replay at least ten KCMO permits and compare actual requirements/outcomes", "Checklist and source quality is credible"],
    ["D. Discovery outreach", "Personalized interviews with 20 activity-qualified companies", "At least five recurring pain confirmations"],
    ["E. Free pilots", "Concierge service for a small number of projects with human review", "Three completed pilots and measurable time saved"],
    ["F. Paid conversion", "Simple per-permit or monthly founding plan", "Two repeat paying customers"],
    ["G. Expansion", "Add adjacent permit type or neighboring jurisdiction based on paid demand", "Repeatability before geographic breadth"]
], [1.15,3.6,2.0])

doc.add_heading("11. Team roles and expert access", level=2)
table(["Owner", "Responsibility"], [
    ["Product/engineering", "Workflow, reliability, evidence provenance, data model, security, analytics"],
    ["Customer discovery lead", "Interview scheduling, notes, pilot expectations, follow-up, outcome measurement"],
    ["Knowledge operations", "Official source collection, update monitoring, rule review, uncertainty tracking"],
    ["Part-time domain advisor", "Reviews workflow and early cases; defines escalation boundaries"],
    ["Customer/licensed applicant", "Confirms facts, licenses, representations, filing, fees, and final approval"],
    ["Architect/engineer when needed", "Licensed technical judgment, signed/sealed plans, code interpretations"]
], [2.1,4.65])

doc.add_heading("12. Decisions to make now", level=2)
for x in [
    "Confirm KCMO as the first jurisdiction; do not mix KCMO and KCK rules.",
    "Choose one first permit workflow after five practitioner interviews, with residential mechanical as the leading hypothesis.",
    "Recruit one paid or advisory permit-domain reviewer before real pilot submissions.",
    "Rename 'Approve for Filing' if the product is not actually controlling a filing workflow; use 'Readiness reviewed' or 'Customer approved packet.'",
    "Adopt source freshness and uncertainty as first-class product fields.",
    "Begin warm discovery now; begin cold contractor outreach when the demo and sample report pass the acceptance gates.",
    "Use the official KCMO dataset to generate the first 20-person outreach list and update it regularly."
]: numbered(x)

page_break()
doc.add_heading("Appendix A - Research sources", level=1)
sources=[
    ("KCMO Permits Division", "https://www.kcmo.gov/city-hall/departments/city-planning-development/permits-division/"),
    ("KCMO Development Process Guide", "https://www.kcmo.gov/city-hall/departments/city-planning-development/development-management/development-guide"),
    ("CompassKC guides", "https://www.kcmo.gov/city-hall/departments/city-planning-development/compass-kc-the-new-permitting-system"),
    ("Electrical, plumbing and mechanical permits", "https://www.kcmo.gov/city-hall/departments/city-planning-development/electrical-plumbing-and-mechanical-permits"),
    ("KCMO Inspections Division", "https://www.kcmo.gov/city-hall/departments/city-planning-development/inspections-division"),
    ("KCMO Information Bulletin index", "https://www.kcmo.gov/city-hall/departments/city-planning-development/information-bulletin-ib-index"),
    ("KCMO code questions", "https://www.kcmo.gov/city-hall/departments/city-planning-development/development-concierge/code-questions"),
    ("KCMO contractor licensing", "https://www.kcmo.gov/city-hall/departments/city-planning-development/contractor-licensing"),
    ("KCMO Permits CPD open dataset", "https://data.kcmo.org/Development/Permits-CPD-Dataset/ntw8-aacc"),
    ("KCMO permit dashboards", "https://data.kcmo.org/stories/s/Building-Permits-Dashboards/sq5v-m7n2/"),
    ("AIA Kansas City", "https://www.aiakc.org/"),
    ("Kansas City Home Builders Association", "https://kchba.org/"),
    ("Kansas City American Subcontractors Association", "https://www.kcasa.org/"),
    ("Kansas PHCC members", "https://www.phccks.org/members/"),
    ("PermitFlow product comparison", "https://www.permitflow.com/permit-management")
]
for name,url in sources:
    para=doc.add_paragraph(style="List Bullet"); r=para.add_run(name+": "); font(r,9.5,True,NAVY)
    r=para.add_run(url); font(r,9.2,False,BLUE)

doc.add_heading("Appendix B - Research limitations", level=1)
for x in [
    "This document is product and market research, not legal, architectural, engineering, or permitting advice.",
    "The official permit dataset's latest identified data date was May 9, 2025; current outreach requires refreshed verification.",
    "Raw permit counts are not company-quality ratings and can contain duplicate business names, DBAs, and related entities.",
    "The authenticated PermitFlow application was not accessed; comparison relies on public product materials.",
    "KCMO, Kansas City Kansas/Wyandotte County, Johnson County municipalities, and St. Louis are separate permitting systems.",
    "Before a live pilot, counsel should review customer authorization, data terms, disclaimers, and professional-services boundaries."
]: bullet(x)

# Keep table rows together where practical and repeat headers
for t in doc.tables:
    # Normalize every table to exact fixed DXA geometry across the 6.75-inch design width.
    target_width = 9720
    tblPr = t._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.insert(0, tblW)
    tblW.set(qn("w:type"), "dxa"); tblW.set(qn("w:w"), str(target_width))
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd"); tblPr.append(tblInd)
    first_mar = t.rows[0].cells[0]._tc.get_or_add_tcPr().first_child_found_in("w:tcMar")
    start_mar = first_mar.find(qn("w:start")) if first_mar is not None else None
    indent = start_mar.get(qn("w:w")) if start_mar is not None else "110"
    tblInd.set(qn("w:type"), "dxa"); tblInd.set(qn("w:w"), indent)
    raw = []
    for c in t.rows[0].cells:
        tcW = c._tc.get_or_add_tcPr().find(qn("w:tcW"))
        raw.append(int(tcW.get(qn("w:w"))) if tcW is not None and tcW.get(qn("w:w")) else 1)
    total = sum(raw) or len(raw)
    widths = [round(target_width * value / total) for value in raw]
    widths[-1] += target_width - sum(widths)
    grid = t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(width)); grid.append(gc)
    for row in t.rows:
        for idx, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None: tcW = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().insert(0, tcW)
            tcW.set(qn("w:type"), "dxa"); tcW.set(qn("w:w"), str(widths[idx]))
    for row in t.rows:
        trPr=row._tr.get_or_add_trPr(); cant=OxmlElement("w:cantSplit"); trPr.append(cant)
    trPr=t.rows[0]._tr.get_or_add_trPr(); hdr=OxmlElement("w:tblHeader"); hdr.set(qn("w:val"),"true"); trPr.append(hdr)

doc.core_properties.title="EstatePermit Kansas City MVP and Go-to-Market Playbook"
doc.core_properties.subject="Product specification, KCMO permitting research, target market, and outreach plan"
doc.core_properties.author="EstatePermit"
doc.save(OUT)
print(OUT)
